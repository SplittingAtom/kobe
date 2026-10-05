"""Dump a workbook: python xlsx_dump.py FILE.xlsx [--max-rows N] [--formulas]"""

import argparse

from openpyxl import load_workbook


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("file")
    ap.add_argument("--max-rows", type=int, default=20)
    ap.add_argument("--formulas", action="store_true")
    args = ap.parse_args()
    wb = load_workbook(args.file, data_only=not args.formulas, read_only=True)
    for ws in wb.worksheets:
        print(f"== sheet {ws.title!r} dimensions={ws.calculate_dimension()} rows={ws.max_row} cols={ws.max_column}")
        for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if n > args.max_rows:
                print("...")
                break
            print(" | ".join("" if v is None else str(v) for v in row))


if __name__ == "__main__":
    main()
