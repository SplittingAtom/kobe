"""CSV -> formatted .xlsx: python csv_to_xlsx.py DATA.csv OUT.xlsx [--sheet NAME] [--total-row]"""

import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORBIDDEN_SHEET_CHARS = set("[]:*?/\\")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("csv")
    ap.add_argument("output")
    ap.add_argument("--sheet", default="Data")
    ap.add_argument("--total-row", action="store_true")
    args = ap.parse_args()
    if not args.output.lower().endswith(".xlsx"):
        sys.exit("output must end in .xlsx")
    if not args.sheet or len(args.sheet) > 31 or FORBIDDEN_SHEET_CHARS & set(args.sheet):
        sys.exit("invalid sheet name")

    df = pd.read_csv(args.csv)
    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    last = len(df) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{last}"
    if args.total_row and len(df):
        ws.cell(row=last + 1, column=1, value="Total").font = Font(bold=True)
        for idx, col in enumerate(df.columns, start=1):
            if pd.api.types.is_numeric_dtype(df[col]) and not pd.api.types.is_bool_dtype(df[col]):
                letter = get_column_letter(idx)
                ws.cell(row=last + 1, column=idx, value=f"=SUM({letter}2:{letter}{last})").font = Font(
                    bold=True
                )
    for idx, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)
    wb.save(args.output)
    print(f"wrote {args.output}: {len(df)} rows, {len(df.columns)} columns")


if __name__ == "__main__":
    main()
